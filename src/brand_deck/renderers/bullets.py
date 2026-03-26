# renderers/bullets.py — 要点列表页渲染器（也是降级目标）
"""Render a bullet points slide. Also serves as the fallback renderer."""

from __future__ import annotations

from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

from brand_deck.theme import BrandTheme
from brand_deck.utils.pptx_helpers import (
    add_blank_slide, add_title, set_slide_bg, add_shape,
    SLIDE_WIDTH,
)


def render_bullets(prs, slide_data: dict, theme: BrandTheme) -> None:
    """Render a bullet points slide."""
    slide = add_blank_slide(prs)
    set_slide_bg(slide, theme.bg_rgb)

    title = slide_data.get("title", "")
    points = slide_data.get("points", [])

    # If no points, try to extract from body/content fields (fallback mode)
    if not points:
        body = slide_data.get("body", "") or slide_data.get("content", "")
        if body:
            points = [line.strip() for line in body.split("\n") if line.strip()]

    if title:
        add_title(slide, title, theme, top=Inches(0.6))

    if not points:
        return

    # Accent bar on left
    add_shape(
        slide,
        left=Inches(0.8),
        top=Inches(2.0),
        width=Inches(0.08),
        height=Inches(len(points) * 0.6 + 0.2),
        fill_color=theme.accent_rgb,
    )

    # Render each bullet point
    bullet_left = Inches(1.2)
    bullet_top_start = Inches(2.1)
    bullet_width = Inches(10.5)

    for i, point in enumerate(points):
        top = bullet_top_start + Inches(i * 0.65)

        txBox = slide.shapes.add_textbox(
            bullet_left, top, bullet_width, Inches(0.55)
        )
        tf = txBox.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        # Bullet prefix
        p.text = f"●  {point}"
        p.alignment = PP_ALIGN.LEFT

        for run in p.runs:
            run.font.name = theme.body_font
            run.font.size = Pt(16)
            run.font.color.rgb = theme.text_dark_rgb
